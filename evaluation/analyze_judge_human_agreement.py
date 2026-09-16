"""
Part I, Stage 2: computes judge-human agreement on the now-completed
evaluation/results/JUDGE_HUMAN_CALIBRATION_40.xlsx, against the k=3 LLM judge
scores already stored in evaluation/results/k_ablation_sweep.json.

Zero new API calls: reads only two already-local files (the filled xlsx and the
cached sweep JSON). Neither is ever written to by this module.

Validation discipline (mirrors golden_set/export_triage_annotations.py and
evaluation/export_retrieval_inspection.py): the read-only context columns
(example_id, tweet_id, customer_text, generated_reply, retrieved evidence) are
RE-DERIVED from k_ablation_sweep.json (via evaluation.build_judge_human_
calibration_workbook.build_rows() -- the exact same function that built the
workbook in the first place) and diffed against what the workbook actually
contains. Any mismatch raises loudly rather than being silently trusted.

SCOPE: this measures judge-human agreement at k=3 ONLY. The k-ablation sweep
(evaluation/K_ABLATION_SWEEP_RESULTS.md) found the judge's own scores move in
different directions across dimensions as k changes (higher Groundedness but
lower Relevance/Tone with retrieval) -- this result must NOT be generalized to
k=0, k=1, or k=5 without a separate calibration pass at those conditions.

Run: python evaluation/analyze_judge_human_agreement.py
"""
import json
import statistics
import sys
from collections import Counter
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from sklearn.metrics import cohen_kappa_score

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import config
from evaluation.build_judge_human_calibration_workbook import (
    CALIBRATION_K,
    K_ABLATION_SWEEP_PATH,
    XLSX_PATH,
    build_rows,
)

AGREEMENT_OUTPUT_PATH = config.EVAL_DIR / "judge_human_agreement.json"

DIMENSIONS = ("relevance", "groundedness", "helpfulness", "tone")
HUMAN_FIELD_BY_DIMENSION = dict(
    relevance="human_relevance", groundedness="human_groundedness",
    helpfulness="human_helpfulness", tone="human_tone",
)
SCORE_COL_BY_DIMENSION = dict(relevance="K", groundedness="L", helpfulness="M", tone="N")

BOOTSTRAP_N_RESAMPLES = 10000
BOOTSTRAP_SEED = 42  # matches project's established seed convention


class AgreementAnalysisError(Exception):
    """Raised on any validation failure. Never caught silently."""


# ---------------------------------------------------------------------------
# Part 1: read + validate the filled workbook against re-derived ground truth
# ---------------------------------------------------------------------------

def read_workbook_rows(xlsx_path=XLSX_PATH):
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb["CALIBRATION_REVIEW"]
    rows = []
    for r in range(2, ws.max_row + 1):
        example_id = ws.cell(row=r, column=1).value
        if example_id is None or str(example_id).strip() == "":
            continue
        retrieved = []
        for i in range(3):
            base = ord("E") + i * 2
            cust_col, reply_col = chr(base), chr(base + 1)
            retrieved.append(dict(
                customer=ws[f"{cust_col}{r}"].value,
                reply=ws[f"{reply_col}{r}"].value,
            ))
        rows.append(dict(
            row_number=r,
            example_id=str(example_id).strip(),
            tweet_id=str(ws.cell(row=r, column=2).value).strip(),
            customer_text=ws.cell(row=r, column=3).value,
            generated_reply=ws.cell(row=r, column=4).value,
            retrieved=retrieved,
            human_relevance=ws["K" + str(r)].value,
            human_groundedness=ws["L" + str(r)].value,
            human_helpfulness=ws["M" + str(r)].value,
            human_tone=ws["N" + str(r)].value,
            human_notes=ws["O" + str(r)].value,
        ))
    return rows


def validate_human_scores(workbook_rows):
    """Every row must have all 4 human score fields present as whole numbers 1-5.
    Reports (does not silently drop/impute) any incomplete or invalid row."""
    problems = []
    for row in workbook_rows:
        for dim in DIMENSIONS:
            field = HUMAN_FIELD_BY_DIMENSION[dim]
            value = row[field]
            if value is None:
                problems.append(f"{row['example_id']}: {field} is blank")
                continue
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                problems.append(f"{row['example_id']}: {field}={value!r} is not numeric")
                continue
            if float(value) != int(value):
                problems.append(f"{row['example_id']}: {field}={value!r} is not a whole number")
                continue
            if not (1 <= int(value) <= 5):
                problems.append(f"{row['example_id']}: {field}={value!r} is out of range 1-5")
    if problems:
        raise AgreementAnalysisError(
            "Human score validation failed for the following row(s)/field(s) -- "
            "refusing to compute agreement on incomplete/invalid data:\n" + "\n".join(problems)
        )
    return True


def _normalize_text(value):
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def validate_context_matches_ground_truth(workbook_rows):
    """Re-derives example_id/tweet_id/customer_text/generated_reply/retrieved
    evidence from evaluation/results/k_ablation_sweep.json (via build_rows(), the
    SAME function that originally built this workbook) and diffs against what the
    workbook actually contains. Raises loudly on any mismatch instead of trusting
    the workbook blindly."""
    ground_truth_rows = build_rows()  # re-reads k_ablation_sweep.json fresh, zero API calls
    ground_truth_by_tweet = {r["tweet_id"]: r for r in ground_truth_rows}

    if len(workbook_rows) != len(ground_truth_rows):
        raise AgreementAnalysisError(
            f"Expected {len(ground_truth_rows)} rows (re-derived), found {len(workbook_rows)} in workbook."
        )

    mismatches = []
    for row in workbook_rows:
        gt = ground_truth_by_tweet.get(row["tweet_id"])
        if gt is None:
            mismatches.append(f"{row['example_id']}: tweet_id={row['tweet_id']} not found in re-derived ground truth.")
            continue

        checks = [
            ("example_id", row["example_id"], gt["example_id"]),
            ("customer_text", _normalize_text(row["customer_text"]), _normalize_text(gt["customer_text"])),
            ("generated_reply", _normalize_text(row["generated_reply"]), _normalize_text(gt["generated_reply"])),
        ]
        for field, actual, expected in checks:
            if actual != expected:
                mismatches.append(
                    f"{row['example_id']}: field {field!r} does not match re-derived ground truth. "
                    f"Expected: {expected!r}. Found: {actual!r}."
                )

        if len(row["retrieved"]) != len(gt["retrieved"]):
            mismatches.append(f"{row['example_id']}: retrieved evidence count mismatch "
                               f"({len(row['retrieved'])} in workbook vs {len(gt['retrieved'])} expected).")
        else:
            for i, (actual_ev, expected_ev) in enumerate(zip(row["retrieved"], gt["retrieved"]), start=1):
                if _normalize_text(actual_ev["customer"]) != _normalize_text(expected_ev["customer"]):
                    mismatches.append(f"{row['example_id']}: retrieved example {i} customer text mismatch.")
                if _normalize_text(actual_ev["reply"]) != _normalize_text(expected_ev["reply"]):
                    mismatches.append(f"{row['example_id']}: retrieved example {i} reply text mismatch.")

    if mismatches:
        raise AgreementAnalysisError(
            "Read-only context fields in the workbook do not match re-derived ground truth:\n"
            + "\n".join(mismatches)
        )
    return True


# ---------------------------------------------------------------------------
# Part 2: pull the corresponding k=3 LLM judge scores
# ---------------------------------------------------------------------------

def load_llm_judge_scores_for_examples(workbook_rows):
    """Returns tweet_id -> {relevance, groundedness, helpfulness, tone, reasoning}
    for the k=3 record of each of the 40 workbook rows. Raises if any are missing
    (re-confirms Stage 1's completeness check rather than trusting it)."""
    with open(K_ABLATION_SWEEP_PATH, encoding="utf-8") as f:
        sweep = json.load(f)
    k3_by_tweet = {r["tweet_id"]: r for r in sweep["records"] if r["k"] == CALIBRATION_K}

    tweet_ids = [row["tweet_id"] for row in workbook_rows]
    missing = [t for t in tweet_ids if t not in k3_by_tweet]
    if missing:
        raise AgreementAnalysisError(f"No k={CALIBRATION_K} LLM judge record for tweet_id(s): {missing}")

    scores = {}
    for t in tweet_ids:
        r = k3_by_tweet[t]
        js = r["judge_scores"]
        for dim in DIMENSIONS:
            if js.get(dim) is None:
                raise AgreementAnalysisError(f"tweet_id={t}: judge_scores.{dim} missing.")
        scores[t] = dict(
            relevance=js["relevance"], groundedness=js["groundedness"],
            helpfulness=js["helpfulness"], tone=js["tone"],
            reasoning=r["judge_reasoning"],
        )
    return scores


# ---------------------------------------------------------------------------
# Part 3: per-dimension agreement statistics
# ---------------------------------------------------------------------------

def bootstrap_kappa_ci(human_scores, llm_scores, n_resamples=BOOTSTRAP_N_RESAMPLES, seed=BOOTSTRAP_SEED):
    """Percentile bootstrap 95% CI for quadratic-weighted Cohen's kappa, resampling
    PAIRS (human_i, llm_i) with replacement -- the standard nonparametric approach
    for an agreement statistic at small n, matching this project's established
    'report an uncertainty interval, not a bare point estimate' discipline
    (evaluation/CALIBRATION_RESULTS.md's Wilson intervals for the same reason)."""
    rng = np.random.RandomState(seed)
    n = len(human_scores)
    human_arr, llm_arr = np.array(human_scores), np.array(llm_scores)
    boot_kappas = []
    for _ in range(n_resamples):
        idx = rng.randint(0, n, size=n)
        h, l = human_arr[idx], llm_arr[idx]
        if len(set(h)) < 2 and len(set(l)) < 2:
            continue  # cohen_kappa_score is undefined (or trivially 1.0/nan) with a single class in both
        boot_kappas.append(cohen_kappa_score(h, l, weights="quadratic", labels=[1, 2, 3, 4, 5]))
    if not boot_kappas:
        # Fully degenerate input (both raters constant across every resample, which
        # can only happen if the ORIGINAL human/llm scores for this dimension are
        # each a single constant value) -- kappa itself is undefined here, so the
        # CI is reported as NaN rather than crashing on an empty percentile call.
        return float("nan"), float("nan"), 0
    lo, hi = np.percentile(boot_kappas, [2.5, 97.5])
    return float(lo), float(hi), len(boot_kappas)


def compute_dimension_stats(human_scores, llm_scores):
    n = len(human_scores)
    diffs = [h - l for h, l in zip(human_scores, llm_scores)]

    exact_agree = sum(1 for d in diffs if d == 0)
    within_1 = sum(1 for d in diffs if abs(d) <= 1)

    kappa = cohen_kappa_score(human_scores, llm_scores, weights="quadratic", labels=[1, 2, 3, 4, 5])
    kappa_ci_lo, kappa_ci_hi, n_valid_resamples = bootstrap_kappa_ci(human_scores, llm_scores)

    mean_signed_diff = statistics.mean(diffs)
    stdev_diff = statistics.stdev(diffs) if n > 1 else 0.0

    confusion = Counter(zip(human_scores, llm_scores))  # (human, llm) -> count

    return dict(
        n=n,
        exact_agreement_count=exact_agree, exact_agreement_rate=exact_agree / n,
        within_1_agreement_count=within_1, within_1_agreement_rate=within_1 / n,
        weighted_kappa=float(kappa), kappa_ci_lo=kappa_ci_lo, kappa_ci_hi=kappa_ci_hi,
        kappa_bootstrap_n_valid=n_valid_resamples,
        mean_signed_diff_human_minus_llm=mean_signed_diff, stdev_diff=stdev_diff,
        confusion_matrix={f"human={h}_llm={l}": c for (h, l), c in sorted(confusion.items())},
    )


# ---------------------------------------------------------------------------
# Part 4: 2+-point disagreements, in full
# ---------------------------------------------------------------------------

def find_large_disagreements(workbook_rows, llm_scores_by_tweet, threshold=2):
    disagreements = []
    for row in workbook_rows:
        llm = llm_scores_by_tweet[row["tweet_id"]]
        per_dim_diffs = {}
        for dim in DIMENSIONS:
            human_val = row[HUMAN_FIELD_BY_DIMENSION[dim]]
            llm_val = llm[dim]
            if abs(human_val - llm_val) >= threshold:
                per_dim_diffs[dim] = dict(human=human_val, llm=llm_val, diff=human_val - llm_val)
        if per_dim_diffs:
            disagreements.append(dict(
                example_id=row["example_id"], tweet_id=row["tweet_id"],
                customer_text=row["customer_text"], generated_reply=row["generated_reply"],
                human_notes=row["human_notes"],
                dimensions=per_dim_diffs,
                llm_reasoning=llm["reasoning"],
            ))
    return disagreements


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    workbook_rows = read_workbook_rows()
    print(f"Read {len(workbook_rows)} rows from {XLSX_PATH}.")

    validate_human_scores(workbook_rows)
    print("All 40 rows have complete, valid (whole number, 1-5) human scores for all 4 dimensions.")

    validate_context_matches_ground_truth(workbook_rows)
    print("Context columns (example_id, tweet_id, customer_text, generated_reply, "
          "retrieved evidence) match re-derived ground truth exactly -- 0 mismatches.")

    llm_scores_by_tweet = load_llm_judge_scores_for_examples(workbook_rows)
    print(f"Loaded k={CALIBRATION_K} LLM judge scores for all {len(llm_scores_by_tweet)} examples "
          f"from {K_ABLATION_SWEEP_PATH} (zero new API calls).")

    per_dimension = {}
    for dim in DIMENSIONS:
        human_scores = [row[HUMAN_FIELD_BY_DIMENSION[dim]] for row in workbook_rows]
        llm_scores = [llm_scores_by_tweet[row["tweet_id"]][dim] for row in workbook_rows]
        per_dimension[dim] = compute_dimension_stats(human_scores, llm_scores)

    print(f"\n{'dimension':<14}{'exact':>8}{'within1':>10}{'kappa':>10}{'kappa_ci':>18}{'mean_diff':>11}")
    for dim in DIMENSIONS:
        s = per_dimension[dim]
        print(f"{dim:<14}{s['exact_agreement_rate']:>8.1%}{s['within_1_agreement_rate']:>10.1%}"
              f"{s['weighted_kappa']:>10.3f}   [{s['kappa_ci_lo']:.3f}, {s['kappa_ci_hi']:.3f}]"
              f"{s['mean_signed_diff_human_minus_llm']:>11.3f}")

    disagreements = find_large_disagreements(workbook_rows, llm_scores_by_tweet, threshold=2)
    print(f"\n{len(disagreements)} example(s) with a 2+-point disagreement on at least one dimension.")

    config.ensure_dirs()
    with open(AGREEMENT_OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(dict(
            n_examples=len(workbook_rows), calibration_k=CALIBRATION_K,
            per_dimension=per_dimension,
            large_disagreements=disagreements,
            scope_limitation=(
                "This agreement measurement is valid for k=3 only. The k-ablation sweep found "
                "the judge's own scores move in different directions across dimensions as k "
                "changes (higher Groundedness but lower Relevance/Tone with retrieval) -- this "
                "result must not be generalized to k=0, k=1, or k=5 without separate calibration."
            ),
        ), f, indent=2)
    print(f"\nWrote {AGREEMENT_OUTPUT_PATH}")
    return per_dimension, disagreements


if __name__ == "__main__":
    main()
