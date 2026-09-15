"""
Exports Faiz's actual retrieval-inspection annotations from the human-authored

    evaluation/results/RETRIEVAL_INSPECTION_20.xlsx

into a NEW canonical file:

    evaluation/results/retrieval_inspection_annotations.csv

GRANULARITY NOTE -- read before touching this file or its output shape:

The original scaffold, evaluation/results/retrieval_inspection_scaffold.csv
(built by evaluation/build_retrieval_inspection_scaffold.py), has 100 rows -- one
per (golden example, retrieved-pair) -- and reserved 5 qualitative judgment columns
on EVERY row, implicitly assuming Faiz would make an independent judgment for each
of the 5 retrieved pairs per example (100 potential judgments).

That is not what actually happened. The human-facing workbook (evaluation/
build_retrieval_inspection_workbook.py) was built one row PER EXAMPLE (20 rows),
with all 5 retrieved pairs shown side-by-side in the same row, and a SINGLE summary
judgment made across all 5 (Retrieval Usefulness / Best Rank / Relevance Problem /
Grounding Value / Overall Observation -- columns Y-AC). Faiz's real annotations are
therefore inherently example-level (20 judgments), not pair-level (100).

This script does NOT force those 20 example-level judgments into the old CSV's
100-row per-pair shape -- duplicating one summary judgment across a golden example's
5 rank-rows would misrepresent it as five independent rank-specific judgments, which
were never made. Instead it writes a NEW file, shaped honestly as one row per
example (20 rows). evaluation/results/retrieval_inspection_scaffold.csv is left
completely untouched by this script -- it remains valid, byte-for-byte, as the raw
per-pair retrieval-EVIDENCE reference data; its 5 always-blank qualitative columns
simply reflect that no judgment was ever made at that granularity, and (per this
finding) never will be.

VALIDATION DISCIPLINE (mirrors golden_set/export_triage_annotations.py): every
read-only/factual field the xlsx carries is RE-DERIVED independently --
candidate_id from golden_set/GOLDEN_200_FINAL.csv, predicted_intent from
evaluation/results/llm_classifier_results.json, and the actual retrieved evidence
from a FRESH evaluation.retrieval.retrieve_top_k() call per example (the same
seed=20 sample used to build the scaffold in the first place) -- and compared
field-by-field against what the xlsx contains. Any mismatch raises
RetrievalInspectionExportError loudly; nothing is silently trusted, "corrected", or
overwritten.

API-SPEND SAFETY: before any retrieve_top_k() call, check_no_api_spend_required()
verifies every one of the 20 query-text embeddings is ALREADY a cache hit (by
computing the exact evaluation.embeddings.DiskCache key and checking the file
exists -- no API call is made just to check). If any is missing, this raises rather
than proceeding, so this script can never spend against a credit-exhausted account.

evaluation/results/RETRIEVAL_INSPECTION_20.xlsx is READ-ONLY input to this script
and is never written to.

Run: python evaluation/export_retrieval_inspection.py
"""
import csv
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import config
from evaluation.build_retrieval_inspection_scaffold import load_golden_rows_raw, select_inspection_sample
from evaluation.build_retrieval_inspection_workbook import (
    BEST_RANK_VALUES,
    GROUNDING_VALUE_VALUES,
    RELEVANCE_PROBLEM_VALUES,
    RETRIEVAL_USEFULNESS_VALUES,
    XLSX_PATH,
    load_candidate_id_by_tweet_id,
    load_predicted_intent_by_tweet_id,
)
from evaluation.retrieval import retrieve_top_k

OUTPUT_CSV = config.EVAL_DIR / "retrieval_inspection_annotations.csv"
SHEET_NAME = "RETRIEVAL_REVIEW"
EXPECTED_RANKS = [1, 2, 3, 4, 5]


class RetrievalInspectionExportError(Exception):
    """Raised on any validation failure. Never caught silently -- callers must let
    it propagate rather than proceed with an unverified export."""


def _normalize_text(value):
    if value is None:
        return ""
    text = str(value)
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def _normalize_best_rank(value):
    """Z (Best Rank) is dropdown-validated against the strings "1".."5"/"NONE", but
    Excel stores a typed/selected numeric-looking entry as an actual number (int or
    float), not a string -- a normal, benign openpyxl/Excel quirk, not tampering.
    Normalizes both representations to a canonical string for comparison/export."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return str(int(value))
    return str(value).strip()


# ---------------------------------------------------------------------------
# API-spend safety check (must pass before any retrieve_top_k call)
# ---------------------------------------------------------------------------

def _embedding_cache_path(text):
    cache_key = json.dumps(dict(model=config.EMBEDDING_MODEL, input=text), sort_keys=True)
    digest = hashlib.sha256(cache_key.encode("utf-8")).hexdigest()
    return config.CACHE_DIR / f"{digest}.json"


def check_no_api_spend_required(customer_texts):
    """Raises RetrievalInspectionExportError if ANY of the given texts' embedding is
    not already cached (i.e. re-deriving would require a live API call). Never calls
    the API itself -- only checks for the cache file's existence."""
    missing = [t for t in customer_texts if not _embedding_cache_path(t).exists()]
    if missing:
        raise RetrievalInspectionExportError(
            f"{len(missing)} of {len(customer_texts)} required query-text embeddings are "
            "NOT already cached -- re-deriving ground truth would require a live API call. "
            "Refusing to proceed (the account is credit-exhausted). "
            f"Missing texts (truncated): {[t[:60] for t in missing[:5]]}"
        )


# ---------------------------------------------------------------------------
# Ground-truth re-derivation (independent of the xlsx and of the scaffold CSV)
# ---------------------------------------------------------------------------

def rederive_ground_truth():
    """Returns (expected_tweet_ids, ground_truth) where ground_truth is
    tweet_id -> {candidate_id, customer_text, predicted_intent, ranked_evidence}.
    ranked_evidence is a list of 5 dicts from a FRESH retrieve_top_k(text, k=5) call
    -- not read from the scaffold CSV or the xlsx. Raises via
    check_no_api_spend_required() before any such call if it would need live spend.
    """
    golden_rows = load_golden_rows_raw()
    expected_tweet_ids = select_inspection_sample(golden_rows)  # seed=20, documented
    by_id = {r["tweet_id"]: r for r in golden_rows}

    customer_texts = [by_id[t]["target_message"] for t in expected_tweet_ids]
    check_no_api_spend_required(customer_texts)

    candidate_id_by_tweet = load_candidate_id_by_tweet_id()
    predicted_intent_by_tweet = load_predicted_intent_by_tweet_id()

    ground_truth = {}
    for tweet_id in expected_tweet_ids:
        customer_text = by_id[tweet_id]["target_message"]
        retrieved = retrieve_top_k(customer_text, k=5)  # fresh, live re-derivation (cache-hit only)
        ranked_evidence = [
            dict(rank=i, similarity=round(sim, 4), hist_customer=cust, hist_brand=brand)
            for i, (cust, brand, sim, _meta) in enumerate(retrieved, start=1)
        ]
        ground_truth[tweet_id] = dict(
            candidate_id=candidate_id_by_tweet[tweet_id],
            customer_text=customer_text,
            predicted_intent=predicted_intent_by_tweet[tweet_id],
            ranked_evidence=ranked_evidence,
        )
    return expected_tweet_ids, ground_truth


# ---------------------------------------------------------------------------
# Read the xlsx (read-only)
# ---------------------------------------------------------------------------

def read_workbook_rows(xlsx_path=XLSX_PATH):
    wb = load_workbook(xlsx_path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise RetrievalInspectionExportError(
            f"Expected sheet {SHEET_NAME!r} not found in {xlsx_path}. Sheets present: {wb.sheetnames}"
        )
    ws = wb[SHEET_NAME]

    rows = []
    for row_idx in range(2, ws.max_row + 1):
        tweet_id_val = ws.cell(row=row_idx, column=2).value
        if tweet_id_val is None or str(tweet_id_val).strip() == "":
            continue

        ranked_evidence = []
        for rank in EXPECTED_RANKS:
            base = ord("E") + (rank - 1) * 4
            rank_col, sim_col, cust_col, brand_col = (chr(base), chr(base + 1), chr(base + 2), chr(base + 3))
            ranked_evidence.append(dict(
                rank=ws[f"{rank_col}{row_idx}"].value,
                similarity=ws[f"{sim_col}{row_idx}"].value,
                hist_customer=ws[f"{cust_col}{row_idx}"].value,
                hist_brand=ws[f"{brand_col}{row_idx}"].value,
            ))

        rows.append(dict(
            row_number=row_idx,
            candidate_id=_normalize_text(ws.cell(row=row_idx, column=1).value),
            tweet_id=str(tweet_id_val).strip(),
            customer_text=ws.cell(row=row_idx, column=3).value,
            predicted_intent=_normalize_text(ws.cell(row=row_idx, column=4).value),
            ranked_evidence=ranked_evidence,
            retrieval_usefulness=_normalize_text(ws["Y" + str(row_idx)].value),
            best_rank=_normalize_best_rank(ws["Z" + str(row_idx)].value),
            relevance_problem=_normalize_text(ws["AA" + str(row_idx)].value),
            grounding_value=_normalize_text(ws["AB" + str(row_idx)].value),
            overall_observation=_normalize_text(ws["AC" + str(row_idx)].value),
        ))
    return rows


# ---------------------------------------------------------------------------
# Validation (factual fields) + extraction (judgment fields)
# ---------------------------------------------------------------------------

def validate_and_extract(workbook_rows, expected_tweet_ids, ground_truth):
    if len(workbook_rows) != len(expected_tweet_ids):
        raise RetrievalInspectionExportError(
            f"Expected exactly {len(expected_tweet_ids)} annotation rows in the workbook, "
            f"found {len(workbook_rows)}."
        )

    actual_ids = {r["tweet_id"] for r in workbook_rows}
    expected_ids_set = set(expected_tweet_ids)
    if actual_ids != expected_ids_set:
        missing = expected_ids_set - actual_ids
        extra = actual_ids - expected_ids_set
        raise RetrievalInspectionExportError(
            "Workbook tweet_ids do not match the re-derived seed=20 sample. "
            f"Missing: {sorted(missing, key=int)}. Unexpected: {sorted(extra, key=int)}."
        )

    mismatches = []
    for r in workbook_rows:
        tid = r["tweet_id"]
        gt = ground_truth[tid]

        checks = [
            ("candidate_id", r["candidate_id"], gt["candidate_id"]),
            ("customer_text", _normalize_text(r["customer_text"]), _normalize_text(gt["customer_text"])),
            ("predicted_intent", r["predicted_intent"], gt["predicted_intent"]),
        ]
        for field, actual, expected in checks:
            if actual != expected:
                mismatches.append(
                    f"tweet_id={tid}: field {field!r} does not match re-derived ground truth. "
                    f"Expected: {expected!r}. Found in workbook: {actual!r}."
                )

        for actual_ev, expected_ev in zip(r["ranked_evidence"], gt["ranked_evidence"]):
            rank = expected_ev["rank"]
            if actual_ev["rank"] != rank:
                mismatches.append(f"tweet_id={tid} rank={rank}: rank cell value {actual_ev['rank']!r} != {rank}.")
            actual_sim = actual_ev["similarity"]
            if actual_sim is None or round(float(actual_sim), 4) != expected_ev["similarity"]:
                mismatches.append(
                    f"tweet_id={tid} rank={rank}: similarity {actual_sim!r} != re-derived {expected_ev['similarity']!r}."
                )
            if _normalize_text(actual_ev["hist_customer"]) != _normalize_text(expected_ev["hist_customer"]):
                mismatches.append(f"tweet_id={tid} rank={rank}: retrieved customer text does not match re-derived retrieval.")
            if _normalize_text(actual_ev["hist_brand"]) != _normalize_text(expected_ev["hist_brand"]):
                mismatches.append(f"tweet_id={tid} rank={rank}: retrieved brand-reply text does not match re-derived retrieval.")

    if mismatches:
        raise RetrievalInspectionExportError(
            "Read-only/factual fields in the workbook do not match independently re-derived "
            "ground truth -- refusing to export. This does not necessarily mean anything was "
            "edited; it could also mean the retrieval index or re-derivation inputs changed. "
            "Investigate before proceeding:\n" + "\n".join(mismatches)
        )

    extracted = []
    for r in workbook_rows:
        tid = r["tweet_id"]
        gt = ground_truth[tid]

        ru = r["retrieval_usefulness"]
        if ru and ru not in RETRIEVAL_USEFULNESS_VALUES:
            raise RetrievalInspectionExportError(
                f"tweet_id={tid}: invalid retrieval_usefulness {ru!r}. Allowed: {RETRIEVAL_USEFULNESS_VALUES} (or blank)."
            )
        br = r["best_rank"]
        if br and br not in BEST_RANK_VALUES:
            raise RetrievalInspectionExportError(
                f"tweet_id={tid}: invalid best_rank {br!r}. Allowed: {BEST_RANK_VALUES} (or blank)."
            )
        rp = r["relevance_problem"]
        if rp and rp not in RELEVANCE_PROBLEM_VALUES:
            raise RetrievalInspectionExportError(
                f"tweet_id={tid}: invalid relevance_problem {rp!r}. Allowed: {RELEVANCE_PROBLEM_VALUES} (or blank)."
            )
        gv = r["grounding_value"]
        if gv and gv not in GROUNDING_VALUE_VALUES:
            raise RetrievalInspectionExportError(
                f"tweet_id={tid}: invalid grounding_value {gv!r}. Allowed: {GROUNDING_VALUE_VALUES} (or blank)."
            )

        extracted.append(dict(
            example_id=r["candidate_id"],
            tweet_id=tid,
            customer_text=gt["customer_text"],
            predicted_intent=gt["predicted_intent"],
            retrieval_usefulness=ru,
            best_rank=br,
            relevance_problem=rp,
            grounding_value=gv,
            overall_observation=r["overall_observation"],
        ))

    # Deterministic output order: sorted by tweet_id ascending (matches the
    # workbook's own build order, but not assumed -- enforced here explicitly).
    extracted.sort(key=lambda r: int(r["tweet_id"]))
    return extracted


HEADER_COMMENT = """\
# retrieval_inspection_annotations.csv
#
# Canonical record of Faiz's manual retrieval-inspection judgments
# (implementation_plan.md \xa79, "Manual retrieval inspection"), exported from
# evaluation/results/RETRIEVAL_INSPECTION_20.xlsx by
# evaluation/export_retrieval_inspection.py.
#
# ONE ROW PER GOLDEN EXAMPLE (20 rows) -- NOT one row per retrieved pair.
#
# GRANULARITY NOTE: the original scaffold, evaluation/results/
# retrieval_inspection_scaffold.csv, has 100 rows (20 examples x 5 retrieved
# pairs) and reserved a qualitative-judgment column on every one of those rows,
# assuming a per-pair judgment. That was never how the annotation was actually
# done: the human-facing workbook presents all 5 retrieved pairs for an example
# side-by-side and asks for ONE summary judgment across all 5. This file reflects
# that real granularity honestly rather than duplicating one judgment across 5
# rows as if they were independent. retrieval_inspection_scaffold.csv is untouched
# and remains valid as raw per-pair retrieval-evidence reference data; its 5
# qualitative columns are simply not used going forward.
#
# example_id, tweet_id, predicted_intent are copied read-only from
# golden_set/GOLDEN_200_FINAL.csv (candidate_id/tweet_id) and evaluation/results/
# llm_classifier_results.json (predicted_intent, Experiment 1's output -- NOT the
# gold intent label). customer_text is the golden example's target_message.
#
# retrieval_usefulness, best_rank, relevance_problem, grounding_value,
# overall_observation are Faiz's own judgments, entered via the xlsx workbook.
# No automated process infers or suggests these values.
#
# Every field above (except the 5 judgment columns) was independently re-derived
# from source data (GOLDEN_200_FINAL.csv, llm_classifier_results.json, and a fresh
# evaluation.retrieval.retrieve_top_k() call) and verified to match the workbook
# before this file was written -- see export_retrieval_inspection.py's
# validate_and_extract().
"""


def write_canonical_csv(extracted_rows, path=OUTPUT_CSV):
    fieldnames = [
        "example_id", "tweet_id", "customer_text", "predicted_intent",
        "retrieval_usefulness", "best_rank", "relevance_problem", "grounding_value",
        "overall_observation",
    ]
    with open(path, "w", newline="", encoding="utf-8") as f:
        f.write(HEADER_COMMENT)
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for r in extracted_rows:
            writer.writerow(r)


def export(xlsx_path=XLSX_PATH, output_csv=OUTPUT_CSV):
    expected_tweet_ids, ground_truth = rederive_ground_truth()
    workbook_rows = read_workbook_rows(xlsx_path)
    extracted = validate_and_extract(workbook_rows, expected_tweet_ids, ground_truth)
    write_canonical_csv(extracted, output_csv)
    return extracted


def main():
    extracted = export()
    n_full = sum(
        1 for r in extracted
        if r["retrieval_usefulness"] and r["best_rank"] and r["relevance_problem"] and r["grounding_value"]
    )
    print(f"Wrote {OUTPUT_CSV} ({len(extracted)} rows, {n_full} with all 4 required judgment fields filled).")


if __name__ == "__main__":
    main()
