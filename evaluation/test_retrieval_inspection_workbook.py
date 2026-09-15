"""
Tests for evaluation/build_retrieval_inspection_workbook.py and its output,
evaluation/results/RETRIEVAL_INSPECTION_20.xlsx.

No API calls are made anywhere in this suite. Everything is checked against the
already-generated workbook and/or a freshly-built isolated copy in a temp dir (never
overwriting the live workbook if it has already been annotated).
"""
import csv
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from evaluation.build_retrieval_inspection_workbook import (
    BEST_RANK_VALUES,
    COLUMNS,
    GROUNDING_VALUE_VALUES,
    RELEVANCE_PROBLEM_VALUES,
    RETRIEVAL_USEFULNESS_VALUES,
    SCAFFOLD_CSV_PATH,
    XLSX_PATH,
    build_examples,
    build_guide_sheet,
    build_reference_sheet,
    build_review_sheet,
    load_candidate_id_by_tweet_id,
    load_predicted_intent_by_tweet_id,
    load_scaffold_rows,
    validate_source_rows,
)


def load_source_csv_rows():
    with open(SCAFFOLD_CSV_PATH, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


@pytest.fixture(scope="module")
def workbook_path():
    assert XLSX_PATH.exists(), f"{XLSX_PATH} does not exist -- run evaluation/build_retrieval_inspection_workbook.py first."
    return XLSX_PATH


@pytest.fixture(scope="module")
def source_rows():
    return load_source_csv_rows()


@pytest.fixture
def fresh_workbook_path(tmp_path):
    """Builds an isolated, freshly-generated workbook in a temp dir using the same
    functions as the real script -- never touches or depends on the live xlsx's
    current (possibly already-annotated) state."""
    rows = load_source_csv_rows()
    examples = build_examples(rows)
    wb = Workbook()
    build_review_sheet(wb, examples)
    build_guide_sheet(wb)
    build_reference_sheet(wb)
    path = tmp_path / "fresh_retrieval_inspection.xlsx"
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Source CSV integrity (read-only) -- validate_source_rows()
# ---------------------------------------------------------------------------

def test_source_csv_has_exactly_20_examples_of_5_ranks_each(source_rows):
    by_example = validate_source_rows(source_rows)
    assert len(by_example) == 20
    for tweet_id, rows in by_example.items():
        assert [int(r["rank"]) for r in rows] == [1, 2, 3, 4, 5]


def test_validate_source_rows_rejects_missing_rank():
    rows = load_source_csv_rows()
    tweet_id = rows[0]["golden_tweet_id"]
    filtered = [r for r in rows if not (r["golden_tweet_id"] == tweet_id and r["rank"] == "3")]
    with pytest.raises(AssertionError, match="expected ranks"):
        validate_source_rows(filtered)


def test_validate_source_rows_rejects_missing_example():
    rows = load_source_csv_rows()
    first_id = rows[0]["golden_tweet_id"]
    filtered = [r for r in rows if r["golden_tweet_id"] != first_id]
    with pytest.raises(AssertionError, match="Expected 20"):
        validate_source_rows(filtered)


def test_validate_source_rows_rejects_empty_required_field():
    rows = [dict(r) for r in load_source_csv_rows()]
    rows[0]["retrieved_brand_text_raw"] = ""
    with pytest.raises(AssertionError, match="empty"):
        validate_source_rows(rows)


# ---------------------------------------------------------------------------
# Workbook structure
# ---------------------------------------------------------------------------

def test_workbook_exists_and_has_expected_sheets(workbook_path):
    wb = load_workbook(workbook_path)
    assert set(wb.sheetnames) == {"RETRIEVAL_REVIEW", "GUIDE", "REFERENCE"}


def test_exactly_20_annotation_rows(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    data_rows = [r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    assert len(data_rows) == 20


def test_header_row_has_expected_column_labels(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    expected = {letter: header for letter, header, *_ in COLUMNS}
    for letter, header in expected.items():
        assert ws[f"{letter}1"].value == header, f"column {letter}: expected {header!r}, got {ws[f'{letter}1'].value!r}"


def test_dropdown_validation_exists_for_y_z_aa_ab(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    dvs = ws.data_validations.dataValidation

    def dv_for(col_letter):
        matches = [dv for dv in dvs if any(str(r).startswith(f"{col_letter}2") for r in dv.sqref.ranges)]
        assert len(matches) == 1, f"expected exactly 1 data validation targeting column {col_letter}, found {len(matches)}"
        return matches[0]

    dv_y = dv_for("Y")
    assert dv_y.type == "list"
    for v in RETRIEVAL_USEFULNESS_VALUES:
        assert v in dv_y.formula1

    dv_z = dv_for("Z")
    for v in BEST_RANK_VALUES:
        assert v in dv_z.formula1

    dv_aa = dv_for("AA")
    for v in RELEVANCE_PROBLEM_VALUES:
        assert v in dv_aa.formula1

    dv_ab = dv_for("AB")
    for v in GROUNDING_VALUE_VALUES:
        assert v in dv_ab.formula1


def test_no_dropdown_validation_on_overall_observation_column():
    # AC must stay free text -- no list validation should target it.
    wb = load_workbook(XLSX_PATH)
    ws = wb["RETRIEVAL_REVIEW"]
    dvs = ws.data_validations.dataValidation
    ac_dvs = [dv for dv in dvs if any(str(r).startswith("AC2") for r in dv.sqref.ranges)]
    assert ac_dvs == []


def test_freeze_panes_header_row_only_no_frozen_columns(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    assert ws.freeze_panes == "A2"
    pane = ws.sheet_view.pane
    assert pane is not None
    assert pane.state == "frozen"
    # No frozen columns -- horizontal scrolling must be unrestricted. openpyxl
    # represents a pure row-only freeze with xSplit unset (None), not 0 -- both
    # mean "no columns frozen"; the object repr itself has no xSplit component.
    assert pane.xSplit in (0, None)
    # Header row stays frozen/visible.
    assert pane.ySplit >= 1


def test_autofilter_covers_full_range(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    assert ws.auto_filter.ref == "A1:AC21"


def test_wrap_text_set_on_message_columns(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    for col in ("C", "G", "H", "K", "L", "O", "P", "S", "T", "W", "X", "AC"):
        assert ws[f"{col}2"].alignment.wrap_text is True, f"column {col} should wrap"
    assert ws["C2"].alignment.vertical == "top"


def test_row_heights_are_reasonable_for_reading(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    for r in range(2, 22):
        assert ws.row_dimensions[r].height >= 45


def test_source_and_retrieved_columns_locked_annotation_columns_unlocked(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    assert ws.protection.sheet is True
    for col in ("A", "B", "C", "D", "E", "F", "G", "H"):
        assert ws[f"{col}2"].protection.locked is True, f"{col} should be locked (source/retrieved)"
    for col in ("Y", "Z", "AA", "AB", "AC"):
        assert ws[f"{col}2"].protection.locked is False, f"{col} should be unlocked (annotation)"


def test_header_comments_present_on_dropdown_columns(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    for col in ("Y", "Z", "AA", "AB"):
        assert ws[f"{col}1"].comment is not None, f"expected a header comment/tooltip on column {col}"


def test_all_annotation_fields_blank_in_freshly_built_workbook(fresh_workbook_path):
    wb = load_workbook(fresh_workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    for r in range(2, 22):
        for col in ("Y", "Z", "AA", "AB", "AC"):
            assert ws[f"{col}{r}"].value in (None, ""), f"{col}{r} should be blank, got {ws[f'{col}{r}'].value!r}"


# ---------------------------------------------------------------------------
# Source-data fidelity: workbook content must exactly match the source CSV
# ---------------------------------------------------------------------------

def test_workbook_customer_messages_and_ids_match_source_csv_exactly(workbook_path, source_rows):
    wb = load_workbook(workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    by_example = validate_source_rows(source_rows)
    candidate_id_by_tweet = load_candidate_id_by_tweet_id()

    for r in range(2, 22):
        tweet_id = str(ws[f"B{r}"].value)
        assert tweet_id in by_example
        first = by_example[tweet_id][0]
        assert ws[f"A{r}"].value == candidate_id_by_tweet[tweet_id]
        assert ws[f"C{r}"].value == first["golden_customer_text"]


def test_workbook_predicted_intents_match_exp1_results_not_gold(workbook_path, source_rows):
    wb = load_workbook(workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    predicted_by_tweet = load_predicted_intent_by_tweet_id()

    for r in range(2, 22):
        tweet_id = str(ws[f"B{r}"].value)
        assert ws[f"D{r}"].value == predicted_by_tweet[tweet_id]


def test_workbook_all_5_retrievals_per_row_preserved_exactly(workbook_path, source_rows):
    wb = load_workbook(workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    by_example = validate_source_rows(source_rows)

    for r in range(2, 22):
        tweet_id = str(ws[f"B{r}"].value)
        source_ranked = by_example[tweet_id]
        for rank in (1, 2, 3, 4, 5):
            base = ord("E") + (rank - 1) * 4
            rank_col, sim_col, cust_col, brand_col = (chr(base), chr(base + 1), chr(base + 2), chr(base + 3))
            source_row = next(sr for sr in source_ranked if int(sr["rank"]) == rank)

            assert ws[f"{rank_col}{r}"].value == rank
            assert ws[f"{sim_col}{r}"].value == pytest.approx(float(source_row["similarity_score"]))
            assert ws[f"{cust_col}{r}"].value == source_row["retrieved_customer_text"]
            assert ws[f"{brand_col}{r}"].value == source_row["retrieved_brand_text_raw"]


def test_workbook_row_order_matches_sorted_tweet_id_ascending(workbook_path, source_rows):
    wb = load_workbook(workbook_path)
    ws = wb["RETRIEVAL_REVIEW"]
    by_example = validate_source_rows(source_rows)
    expected_order = sorted(by_example.keys(), key=int)
    actual_order = [str(ws[f"B{r}"].value) for r in range(2, 22)]
    assert actual_order == expected_order


# ---------------------------------------------------------------------------
# Determinism / no source mutation
# ---------------------------------------------------------------------------

def test_builder_is_deterministic_across_two_runs():
    rows1 = load_scaffold_rows()
    rows2 = load_scaffold_rows()
    examples1 = build_examples(rows1)
    examples2 = build_examples(rows2)
    assert examples1 == examples2


def test_source_csv_is_never_written_by_the_builder():
    before = SCAFFOLD_CSV_PATH.stat().st_mtime
    load_scaffold_rows()
    build_examples(load_scaffold_rows())
    after = SCAFFOLD_CSV_PATH.stat().st_mtime
    assert before == after


def test_no_network_or_api_key_dependency_in_builder_module():
    """The builder module must not IMPORT or CALL the OpenAI client / api-key
    plumbing -- this is a pure CSV/JSON -> xlsx presentation layer. (Mentioning
    "OpenAI API" in a docstring/comment, e.g. to say a call is NOT made, is fine;
    only actual import/call sites are checked here.)"""
    import evaluation.build_retrieval_inspection_workbook as mod
    assert not hasattr(mod, "OpenAI")
    assert not hasattr(mod, "openai")
    src = Path(mod.__file__).read_text(encoding="utf-8")
    assert "import openai" not in src
    assert "from openai" not in src
    assert "get_api_key()" not in src
    assert "client.embeddings" not in src
    assert "client.chat" not in src
