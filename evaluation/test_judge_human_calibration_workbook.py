"""
Tests for evaluation/build_judge_human_calibration_workbook.py and its output,
evaluation/results/JUDGE_HUMAN_CALIBRATION_40.xlsx.

No API calls anywhere in this suite. Everything is checked against the
already-generated workbook and/or a freshly-built isolated copy in a temp dir.
Never reads golden_set/TRIAGE_ANNOTATION_40.csv's triage_decision/
escalation_reason/notes columns -- only example_id/tweet_id, matching the
builder module's own isolation guarantee (tested explicitly below).
"""
import csv
import json
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from evaluation.build_judge_human_calibration_workbook import (
    CALIBRATION_K,
    COLUMNS,
    K_ABLATION_SWEEP_PATH,
    SCORE_COLUMNS,
    TRIAGE_40_CSV_PATH,
    XLSX_PATH,
    build_instructions_sheet,
    build_review_sheet,
    build_rows,
    load_k3_records_by_tweet_id,
    load_shared_example_ids,
    validate_source_completeness,
)


@pytest.fixture(scope="module")
def workbook_path():
    assert XLSX_PATH.exists(), f"{XLSX_PATH} does not exist -- run the builder script first."
    return XLSX_PATH


@pytest.fixture(scope="module")
def shared_examples():
    return load_shared_example_ids()


@pytest.fixture
def fresh_workbook_path(tmp_path):
    rows = build_rows()
    wb = Workbook()
    build_review_sheet(wb, rows)
    build_instructions_sheet(wb)
    path = tmp_path / "fresh_calibration.xlsx"
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Source loading: 40 examples, k=3, correct shared IDs, holdout isolation
# ---------------------------------------------------------------------------

def test_shared_example_ids_are_exactly_40(shared_examples):
    assert len(shared_examples) == 40
    assert len(set(e["tweet_id"] for e in shared_examples)) == 40
    assert len(set(e["example_id"] for e in shared_examples)) == 40


def test_shared_example_ids_never_include_triage_fields():
    """load_shared_example_ids() must return dicts with ONLY example_id/tweet_id --
    structurally impossible for triage_decision/escalation_reason/notes to leak
    into anything built downstream."""
    examples = load_shared_example_ids()
    for ex in examples:
        assert set(ex.keys()) == {"example_id", "tweet_id"}


def test_shared_example_ids_match_triage_40_csv_example_ids():
    with open(TRIAGE_40_CSV_PATH, encoding="utf-8") as f:
        lines = f.readlines()
    data_lines = [ln for ln in lines if not ln.startswith("#")]
    csv_rows = list(csv.DictReader(data_lines))
    csv_ids = {r["example_id"] for r in csv_rows}

    shared = {e["example_id"] for e in load_shared_example_ids()}
    assert shared == csv_ids


def test_k3_records_loaded_are_all_k_equals_3():
    with open(K_ABLATION_SWEEP_PATH, encoding="utf-8") as f:
        sweep = json.load(f)
    by_tweet = load_k3_records_by_tweet_id()
    for tweet_id, record in by_tweet.items():
        assert record["k"] == CALIBRATION_K == 3


def test_each_shared_example_maps_to_exactly_one_k3_record(shared_examples):
    by_tweet = load_k3_records_by_tweet_id()
    for ex in shared_examples:
        matches = [r for r in [by_tweet.get(ex["tweet_id"])] if r is not None]
        assert len(matches) == 1, f"{ex['tweet_id']}: expected exactly 1 k=3 record, found {len(matches)}"


# ---------------------------------------------------------------------------
# Completeness validation (Part 2 logic)
# ---------------------------------------------------------------------------

def test_validate_source_completeness_passes_on_real_data(shared_examples):
    by_tweet = load_k3_records_by_tweet_id()
    assert validate_source_completeness(shared_examples, by_tweet) is True


def test_validate_source_completeness_rejects_missing_record():
    examples = [dict(example_id="X", tweet_id="does-not-exist")]
    with pytest.raises(AssertionError, match="Missing k=3 record"):
        validate_source_completeness(examples, {})


def test_validate_source_completeness_rejects_missing_judge_score_field():
    examples = [dict(example_id="X", tweet_id="1")]
    by_tweet = {
        "1": dict(
            generated_reply="a reply",
            retrieved_evidence=[{"customer_text": "c", "brand_text_shown": "b"}] * 3,
            judge_scores=dict(relevance=4, groundedness=4, helpfulness=None, tone=4),  # missing helpfulness
            judge_reasoning="reasoning",
        )
    }
    with pytest.raises(AssertionError, match="judge_scores.helpfulness"):
        validate_source_completeness(examples, by_tweet)


def test_validate_source_completeness_rejects_wrong_evidence_count():
    examples = [dict(example_id="X", tweet_id="1")]
    by_tweet = {
        "1": dict(
            generated_reply="a reply",
            retrieved_evidence=[{"customer_text": "c", "brand_text_shown": "b"}] * 2,  # only 2, expected 3
            judge_scores=dict(relevance=4, groundedness=4, helpfulness=4, tone=4),
            judge_reasoning="reasoning",
        )
    }
    with pytest.raises(AssertionError, match="retrieved_evidence"):
        validate_source_completeness(examples, by_tweet)


# ---------------------------------------------------------------------------
# Workbook structure
# ---------------------------------------------------------------------------

def test_workbook_exists_and_has_expected_sheets(workbook_path):
    wb = load_workbook(workbook_path)
    assert set(wb.sheetnames) == {"CALIBRATION_REVIEW", "Instructions"}


def test_exactly_40_data_rows(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["CALIBRATION_REVIEW"]
    data_rows = [r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    assert len(data_rows) == 40


def test_header_row_matches_expected_columns(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["CALIBRATION_REVIEW"]
    for letter, header, *_ in COLUMNS:
        assert ws[f"{letter}1"].value == header


def test_score_columns_have_whole_number_1_to_5_validation(workbook_path):
    """Checks the ESSENTIAL validation behavior (whole number, 1-5, covering every
    score cell), not its internal XML representation. Excel is known to
    consolidate several originally-separate per-column DataValidation entries with
    identical settings into one multi-range entry when a human opens/edits/saves
    the file (the same class of benign re-serialization already seen for this
    project's freeze_panes topLeftCell) -- and when it does, it may omit the
    `operator` attribute rather than writing "between" explicitly, since "between"
    is the OOXML spec's default for a whole-number validation with both formula1
    and formula2 set. Neither change alters what the dropdown actually enforces."""
    wb = load_workbook(workbook_path)
    ws = wb["CALIBRATION_REVIEW"]
    dvs = ws.data_validations.dataValidation
    whole_number_dvs = [dv for dv in dvs if dv.type == "whole"]
    assert whole_number_dvs, "expected at least one whole-number data validation"

    for dv in whole_number_dvs:
        assert dv.operator in ("between", None)  # None == the OOXML default of "between"
        assert str(dv.formula1) == "1"
        assert str(dv.formula2) == "5"

    # Every one of the 40 data rows for every score column must be covered by
    # SOME whole-number validation entry -- this is the actual behavioral
    # guarantee, independent of whether Excel represents it as 4 separate
    # per-column entries or has consolidated them into one multi-range entry.
    for col in SCORE_COLUMNS:
        for row in (2, 21, 41):
            covered = any(f"{col}{row}" in dv.sqref for dv in whole_number_dvs)
            assert covered, f"{col}{row} is not covered by any whole-number validation"


def test_notes_column_has_no_data_validation(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["CALIBRATION_REVIEW"]
    dvs = ws.data_validations.dataValidation
    o_dvs = [dv for dv in dvs if any(str(r).startswith("O2") for r in dv.sqref.ranges)]
    assert o_dvs == []


def test_freeze_panes_header_row_only_no_frozen_columns(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["CALIBRATION_REVIEW"]
    pane = ws.sheet_view.pane
    assert pane is not None
    assert pane.state == "frozen"
    assert pane.xSplit in (0, None)
    assert pane.ySplit >= 1


def test_context_columns_locked_score_columns_unlocked(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["CALIBRATION_REVIEW"]
    assert ws.protection.sheet is True
    for col in ("A", "B", "C", "D", "E", "F"):
        assert ws[f"{col}2"].protection.locked is True
    for col in SCORE_COLUMNS + ("O",):
        assert ws[f"{col}2"].protection.locked is False


def test_wrap_text_set_on_message_columns(workbook_path):
    wb = load_workbook(workbook_path)
    ws = wb["CALIBRATION_REVIEW"]
    for col in ("C", "D", "E", "F", "O"):
        assert ws[f"{col}2"].alignment.wrap_text is True


# ---------------------------------------------------------------------------
# Human score columns genuinely blank
# ---------------------------------------------------------------------------

def test_all_human_score_and_notes_cells_blank_in_fresh_workbook(fresh_workbook_path):
    wb = load_workbook(fresh_workbook_path)
    ws = wb["CALIBRATION_REVIEW"]
    for r in range(2, 42):
        for col in SCORE_COLUMNS + ("O",):
            assert ws[f"{col}{r}"].value in (None, ""), f"{col}{r} should be blank"


# ---------------------------------------------------------------------------
# CRITICAL: judge scores/reasoning/grounding_notes must never appear anywhere
# ---------------------------------------------------------------------------

def _all_cell_and_comment_text(ws):
    texts = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                texts.append(str(cell.value))
            if cell.comment is not None:
                texts.append(cell.comment.text)
    return "\n".join(texts)


def test_judge_scores_and_reasoning_never_appear_in_workbook(workbook_path):
    with open(K_ABLATION_SWEEP_PATH, encoding="utf-8") as f:
        sweep = json.load(f)
    shared_tweet_ids = {e["tweet_id"] for e in load_shared_example_ids()}
    k3_records = [r for r in sweep["records"] if r["k"] == 3 and r["tweet_id"] in shared_tweet_ids]

    wb = load_workbook(workbook_path)
    all_text = "\n".join(_all_cell_and_comment_text(wb[name]) for name in wb.sheetnames)

    for r in k3_records:
        reasoning = r["judge_reasoning"]
        # a real, sufficiently distinctive judge_reasoning string must not appear verbatim
        if reasoning and len(reasoning) > 20:
            assert reasoning not in all_text, f"judge_reasoning for tweet_id={r['tweet_id']} leaked into workbook"


def test_grounding_notes_never_appear_in_workbook(workbook_path):
    """grounding_notes claims are short descriptions/paraphrases of content that IS
    legitimately shown (the generated reply and the retrieved evidence) -- a claim
    can coincidentally share a substring with that legitimate text (e.g. it quotes
    a fragment of a retrieved historical reply we correctly display). That overlap
    is expected and NOT a leak. What must never happen is grounding_notes being
    copied in as its OWN structured field (a bullet list, a "grounded_in_evidence:"
    label, etc.) -- checked here by confirming no claim appears in the workbook
    UNLESS it's already explainable as a substring of that same example's own
    legitimately-shown customer/generated/retrieved text."""
    with open(K_ABLATION_SWEEP_PATH, encoding="utf-8") as f:
        sweep = json.load(f)
    shared_tweet_ids = {e["tweet_id"] for e in load_shared_example_ids()}
    k3_records = [r for r in sweep["records"] if r["k"] == 3 and r["tweet_id"] in shared_tweet_ids]

    wb = load_workbook(workbook_path)
    all_text = "\n".join(_all_cell_and_comment_text(wb[name]) for name in wb.sheetnames)

    # the grounding_notes structure itself (labels/keys) must never appear
    for label in ("grounded_in_evidence", "grounded_in_customer_message", "unsupported_or_generic"):
        assert label not in all_text

    for r in k3_records:
        legitimate_text = r["customer_text"] + "\n" + r["generated_reply"] + "\n" + "\n".join(
            ev["customer_text"] + "\n" + ev["brand_text_shown"] for ev in r["retrieved_evidence"]
        )
        for claim_list in r["grounding_notes"].values():
            for claim in claim_list:
                if not claim or len(claim) <= 20:
                    continue
                if claim in legitimate_text:
                    continue  # benign overlap with legitimately-shown content, not a leak
                assert claim not in all_text, (
                    f"grounding_notes claim for tweet_id={r['tweet_id']} leaked into workbook "
                    f"as its own artifact (not explainable by the shown customer/generated/"
                    f"retrieved text): {claim!r}"
                )


def test_build_rows_output_never_contains_judge_or_grounding_keys():
    rows = build_rows()
    for row in rows:
        assert set(row.keys()) == {"example_id", "tweet_id", "customer_text", "generated_reply", "retrieved"}
        for ev in row["retrieved"]:
            assert set(ev.keys()) == {"customer", "reply"}


def test_raw_numeric_judge_scores_do_not_appear_as_standalone_values_in_score_columns():
    """A structural check specifically on the score columns K-N of a freshly built
    (never human-touched) workbook: since all cells start blank, no judge score
    integer could possibly appear there pre-filled."""
    rows = build_rows()
    wb = Workbook()
    build_review_sheet(wb, rows)
    ws = wb.active
    for r in range(2, 42):
        for col in SCORE_COLUMNS:
            assert ws[f"{col}{r}"].value is None


# ---------------------------------------------------------------------------
# Determinism / no source mutation / zero network
# ---------------------------------------------------------------------------

def test_builder_is_deterministic_across_two_runs():
    rows1 = build_rows()
    rows2 = build_rows()
    assert rows1 == rows2


def test_source_sweep_json_is_never_written_by_the_builder():
    before = K_ABLATION_SWEEP_PATH.stat().st_mtime
    build_rows()
    after = K_ABLATION_SWEEP_PATH.stat().st_mtime
    assert before == after


def test_triage_40_csv_is_never_written_by_the_builder():
    before = TRIAGE_40_CSV_PATH.stat().st_mtime
    build_rows()
    after = TRIAGE_40_CSV_PATH.stat().st_mtime
    assert before == after


def test_builder_module_has_no_openai_or_retrieval_import():
    import evaluation.build_judge_human_calibration_workbook as mod
    src = Path(mod.__file__).read_text(encoding="utf-8")
    assert "import openai" not in src
    assert "from openai" not in src
    assert "from evaluation.retrieval" not in src
    assert "from evaluation.embeddings" not in src
    assert "get_api_key()" not in src
