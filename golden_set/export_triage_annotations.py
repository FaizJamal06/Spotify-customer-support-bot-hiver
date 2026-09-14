"""
Converts golden_set/TRIAGE_ANNOTATION_40.xlsx (the human-facing annotation
workbook) back into golden_set/TRIAGE_ANNOTATION_40.csv (the canonical,
machine-readable annotation output).

This is the counterpart to golden_set/build_triage_workbook.py. The xlsx is
where a human annotates; this script is how those annotations become the
canonical record. It is strict on purpose:

- It validates all 40 examples are present, and in the expected order
  (matching the sampling manifest in golden_set/TRIAGE_ANNOTATION_PROTOCOL.md).
- It validates that the read-only context fields (example_id, tweet_id,
  human_gold_label, target_message) were NOT altered in the workbook, by
  comparing them against golden_set/GOLDEN_200_FINAL.csv (the frozen
  source of truth) -- not against the xlsx's own prior state, so tampering
  cannot self-certify. Any mismatch raises loudly; nothing is silently
  overwritten or "corrected."
- It validates triage_decision, when present, is one of the two allowed
  values (AUTO_HANDLE / HUMAN_ESCALATION). escalation_reason has no
  controlled vocabulary (per golden_set/TRIAGE_ANNOTATION_PROTOCOL.md, which
  defines none), so it is preserved as free text, unvalidated beyond being a
  string.
- Blank triage_decision / escalation_reason / notes are valid and preserved
  as blank -- annotation may be partially complete.

Run: python golden_set/export_triage_annotations.py
"""
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

GOLDEN_SET_DIR = Path(__file__).resolve().parent
GOLDEN_CSV = GOLDEN_SET_DIR / "GOLDEN_200_FINAL.csv"
XLSX_PATH = GOLDEN_SET_DIR / "TRIAGE_ANNOTATION_40.xlsx"
OUTPUT_CSV = GOLDEN_SET_DIR / "TRIAGE_ANNOTATION_40.csv"
PROTOCOL_MD = GOLDEN_SET_DIR / "TRIAGE_ANNOTATION_PROTOCOL.md"

TRIAGE_DECISION_VALUES = {"AUTO_HANDLE", "HUMAN_ESCALATION"}

SHEET_NAME = "TRIAGE_ANNOTATION"
COLUMNS = ["example_id", "tweet_id", "human_gold_label", "target_message",
           "triage_decision", "escalation_reason", "notes"]

HEADER_COMMENT = """\
# TRIAGE_ANNOTATION_40.csv
#
# This is a SEPARATE, ADDITIVE annotation layer on top of the frozen intent
# gold set (golden_set/GOLDEN_200_FINAL.csv). It does NOT modify that file
# or any of its 200 human_gold_label values.
#
# example_id, tweet_id, target_message, human_gold_label are copied
# read-only from GOLDEN_200_FINAL.csv, shown here per explicit approval --
# this is an additive triage layer on top of the existing intent gold, not
# a re-judgment of intent.
#
# triage_decision and escalation_reason are filled in by the human annotator
# (Faiz) via golden_set/TRIAGE_ANNOTATION_40.xlsx and exported to this file by
# golden_set/export_triage_annotations.py. No automated process infers or
# suggests these values.
#
# notes is optional free text, for the annotator's own use.
#
# Sampling method, seed, and the full example_id list are documented in
# golden_set/TRIAGE_ANNOTATION_PROTOCOL.md.
"""


class TriageExportError(Exception):
    """Raised when the workbook fails identity/validation checks. Never
    caught silently by this module -- callers should let it propagate."""


def _normalize_text(value):
    if value is None:
        return ""
    text = str(value)
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def load_expected_ids():
    """The canonical, ordered list of 40 example_ids, parsed out of
    TRIAGE_ANNOTATION_PROTOCOL.md's fenced id block -- kept as a single
    source of truth so the workbook and the protocol can't silently drift
    apart."""
    text = PROTOCOL_MD.read_text(encoding="utf-8")
    start = text.index("## The 40 selected example_ids")
    fence_start = text.index("```", start) + 3
    fence_end = text.index("```", fence_start)
    block = text[fence_start:fence_end]
    ids = [tok.strip().rstrip(",") for tok in block.split() if tok.strip()]
    return [i.rstrip(",") for i in ids]


def load_golden_lookup():
    with open(GOLDEN_CSV, encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    return {r["candidate_id"]: r for r in rows}


def read_workbook_rows(xlsx_path=XLSX_PATH):
    wb = load_workbook(xlsx_path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise TriageExportError(
            f"Expected sheet {SHEET_NAME!r} not found in {xlsx_path}. "
            f"Sheets present: {wb.sheetnames}"
        )
    ws = wb[SHEET_NAME]

    rows = []
    for row_idx in range(2, ws.max_row + 1):
        example_id = ws.cell(row=row_idx, column=1).value
        if example_id is None or str(example_id).strip() == "":
            continue
        rows.append(dict(
            row_number=row_idx,
            example_id=str(example_id).strip(),
            tweet_id=str(ws.cell(row=row_idx, column=2).value or "").strip(),
            human_gold_label=str(ws.cell(row=row_idx, column=3).value or "").strip(),
            target_message=ws.cell(row=row_idx, column=4).value,
            triage_decision=ws.cell(row=row_idx, column=5).value,
            escalation_reason=ws.cell(row=row_idx, column=6).value,
            notes=ws.cell(row=row_idx, column=7).value,
        ))
    return rows


def validate_and_extract(workbook_rows, golden_lookup, expected_ids):
    if len(workbook_rows) != 40:
        raise TriageExportError(
            f"Expected exactly 40 annotation rows in the workbook, found {len(workbook_rows)}."
        )

    actual_ids = [r["example_id"] for r in workbook_rows]
    if actual_ids != expected_ids:
        missing = set(expected_ids) - set(actual_ids)
        extra = set(actual_ids) - set(expected_ids)
        order_note = ""
        if not missing and not extra:
            order_note = " (same 40 ids present, but order does not match the protocol manifest)"
        raise TriageExportError(
            "Workbook example_ids do not match the expected 40-example manifest from "
            f"TRIAGE_ANNOTATION_PROTOCOL.md{order_note}. Missing: {sorted(missing)}. "
            f"Unexpected: {sorted(extra)}."
        )

    mismatches = []
    for r in workbook_rows:
        eid = r["example_id"]
        golden = golden_lookup.get(eid)
        if golden is None:
            mismatches.append(f"{eid}: not found in GOLDEN_200_FINAL.csv at all.")
            continue

        checks = [
            ("tweet_id", r["tweet_id"], golden["tweet_id"]),
            ("human_gold_label", r["human_gold_label"], golden["human_gold_label"]),
            ("target_message", _normalize_text(r["target_message"]), _normalize_text(golden["target_message"])),
        ]
        for field, actual, expected in checks:
            if actual != expected:
                mismatches.append(
                    f"{eid}: field {field!r} was modified in the workbook. "
                    f"Expected (from GOLDEN_200_FINAL.csv): {expected!r}. Found: {actual!r}."
                )

    if mismatches:
        raise TriageExportError(
            "Read-only context fields were modified in the workbook -- refusing to export. "
            "Fix these rows in the workbook (restore the original context values) and re-run:\n"
            + "\n".join(mismatches)
        )

    extracted = []
    for r in workbook_rows:
        triage_decision = _normalize_text(r["triage_decision"])
        if triage_decision and triage_decision not in TRIAGE_DECISION_VALUES:
            raise TriageExportError(
                f"{r['example_id']}: invalid triage_decision {triage_decision!r}. "
                f"Allowed values: {sorted(TRIAGE_DECISION_VALUES)} (or blank)."
            )
        extracted.append(dict(
            example_id=r["example_id"],
            tweet_id=r["tweet_id"],
            target_message=golden_lookup[r["example_id"]]["target_message"],
            human_gold_label=r["human_gold_label"],
            triage_decision=triage_decision,
            escalation_reason=_normalize_text(r["escalation_reason"]),
            notes=_normalize_text(r["notes"]),
        ))
    return extracted


def write_canonical_csv(extracted_rows, path=OUTPUT_CSV):
    with open(path, "w", newline="", encoding="utf-8") as f:
        f.write(HEADER_COMMENT)
        writer = csv.writer(f)
        writer.writerow(["example_id", "tweet_id", "target_message", "human_gold_label",
                          "triage_decision", "escalation_reason", "notes"])
        for r in extracted_rows:
            writer.writerow([
                r["example_id"], r["tweet_id"], r["target_message"], r["human_gold_label"],
                r["triage_decision"], r["escalation_reason"], r["notes"],
            ])


def export(xlsx_path=XLSX_PATH, output_csv=OUTPUT_CSV):
    expected_ids = load_expected_ids()
    golden_lookup = load_golden_lookup()
    workbook_rows = read_workbook_rows(xlsx_path)
    extracted = validate_and_extract(workbook_rows, golden_lookup, expected_ids)
    write_canonical_csv(extracted, output_csv)
    return extracted


def main():
    extracted = export()
    n_filled = sum(1 for r in extracted if r["triage_decision"])
    print(f"Wrote {OUTPUT_CSV} ({len(extracted)} rows, {n_filled} with a triage_decision filled in).")


if __name__ == "__main__":
    main()
